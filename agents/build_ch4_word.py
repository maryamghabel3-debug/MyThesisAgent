# -*- coding: utf-8 -*-
"""ساخت سند Word فصل چهارم (نسخهٔ نمایشی دادهٔ شبیه‌سازی‌شده) — مأموریت ۴۳.

خروجی: output/final/chapter4_simulated.docx
همان قالب فصل‌های ۱-۳: B Nazanin/Times New Roman، راست‌به‌چپ، تیترهای
راست‌چین با bidi معتبر، جداول APA سه‌خطی، شکل با عنوان زیرین وسط‌چین.
"""

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
from build_word import (FA_SIZE, EN_SIZE, TABLE_FONT_PT, LINE_SPACING_PT,
                        add_rich_text, base_paragraph, clean_styles,
                        _reorder_ppr, _set_paragraph_bidi, _set_section_rtl,
                        render_table, setup_heading_styles, setup_section)

ROOT = Path(__file__).resolve().parents[1]
MD_PATH = ROOT / 'output' / 'drafts' / 'chapter4_simulated.md'
FIG_PATH = ROOT / 'output' / 'drafts' / 'ch4_fig1_path_model.png'
OUT_PATH = ROOT / 'output' / 'final' / 'chapter4_simulated.docx'


def add_md_rich(par, text, fa_pt=FA_SIZE, en_pt=EN_SIZE):
    """متن مارک‌داون با بخش‌های **بولد** را با ران‌های فارسی/لاتین می‌افزاید."""
    pos = 0
    for m in re.finditer(r'\*\*(.+?)\*\*', text):
        if m.start() > pos:
            add_rich_text(par, text[pos:m.start()], fa_pt, en_pt)
        add_rich_text(par, m.group(1), fa_pt, en_pt, bold_all=True)
        pos = m.end()
    if pos < len(text):
        add_rich_text(par, text[pos:], fa_pt, en_pt)


def apa_table_borders(table):
    """قالب APA: فقط خط بالای جدول، خط زیر سرستون و خط پایین جدول."""
    tblPr = table._tbl.tblPr
    for old in tblPr.findall(qn('w:tblBorders')):
        tblPr.remove(old)
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'bottom'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), '12')
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), '000000')
        borders.append(el)
    tblPr.append(borders)
    for cell in table.rows[0].cells:
        tcPr = cell._tc.get_or_add_tcPr()
        for old in tcPr.findall(qn('w:tcBorders')):
            tcPr.remove(old)
        cb = OxmlElement('w:tcBorders')
        b = OxmlElement('w:bottom')
        b.set(qn('w:val'), 'single')
        b.set(qn('w:sz'), '12')
        b.set(qn('w:space'), '0')
        b.set(qn('w:color'), '000000')
        cb.append(b)
        tcPr.append(cb)


def render_apa_table(doc, rows):
    """جدول APA با bidi و بدون خطوط عمودی."""
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    # نکته: ران‌های ساده (نه add_rich_text)؛ ترکیب ران‌های دوزبانهٔ صریح با
    # جدول‌های بیش از دو سطر در LibreOffice بارگذاری را مختل می‌کند.
    for r, row in enumerate(rows):
        for c, cell_text in enumerate(row):
            cell = table.cell(r, c)
            cell.text = ''
            par = cell.paragraphs[0]
            _set_paragraph_bidi(par)
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = par.add_run(cell_text.strip())
            run.bold = (r == 0)
            run.font.size = Pt(11)
    apa_table_borders(table)
    # bidiVisual پس از tblBorders الحاق شود (ترتیب تحمل‌شده توسط LibreOffice)
    if table._tbl.tblPr.find(qn('w:bidiVisual')) is None:
        table._tbl.tblPr.append(table._tbl.makeelement(qn('w:bidiVisual'), {}))


def build(md_path=MD_PATH, fig_path=FIG_PATH, out_path=OUT_PATH):
    doc = Document()
    clean_styles(doc)
    setup_heading_styles(doc)
    setup_section(doc.sections[0])
    _set_section_rtl(doc.sections[0])

    lines = Path(md_path).read_text(encoding='utf-8').splitlines()
    i = 0
    pending_fig_caption = None
    while i < len(lines):
        line = lines[i].strip()
        i += 1
        if not line:
            continue
        if line.startswith('# '):
            par = base_paragraph(doc, WD_ALIGN_PARAGRAPH.CENTER, 6, 12)
            add_md_rich(par, '**' + line[2:] + '**', fa_pt=16, en_pt=15)
            continue
        if line.startswith('> '):
            par = base_paragraph(doc, space_after=6)
            add_md_rich(par, line[2:])
            continue
        if line.startswith('@subtitle '):
            par = base_paragraph(doc, WD_ALIGN_PARAGRAPH.CENTER, 0, 10)
            add_md_rich(par, '**' + line[len('@subtitle '):] + '**', fa_pt=14, en_pt=13)
            continue
        if line.startswith('@pagebreak'):
            doc.add_page_break()
            continue
        if line.startswith('## '):
            par = doc.add_paragraph(style='Heading 2')
            _set_paragraph_bidi(par)
            par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            _reorder_ppr(par._p.get_or_add_pPr())
            add_rich_text(par, line[3:], fa_pt=14, en_pt=14, bold_all=True)
            continue
        if line.startswith('|'):
            rows = []
            cur = line
            while cur.startswith('|'):
                cells = [c.strip() for c in cur.strip('|').split('|')]
                if not set(''.join(cells)) <= set('-: '):
                    rows.append(cells)
                if i > len(lines) - 1:
                    break
                cur = lines[i].strip()
                i += 1
            render_apa_table(doc, rows)
            continue
        if line.startswith('!['):
            doc.add_picture(str(fig_path), width=Cm(14))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            if pending_fig_caption:
                par = base_paragraph(doc, WD_ALIGN_PARAGRAPH.CENTER, 3, 6)
                add_md_rich(par, pending_fig_caption, fa_pt=TABLE_FONT_PT, en_pt=10)
                pending_fig_caption = None
            continue
        if line.startswith('**شکل'):
            pending_fig_caption = line
            continue
        if line.startswith('**جدول'):
            par = base_paragraph(doc, WD_ALIGN_PARAGRAPH.CENTER, 9, 2)
            add_md_rich(par, line, fa_pt=TABLE_FONT_PT, en_pt=10)
            continue
        par = base_paragraph(doc, space_after=0)
        add_md_rich(par, line)

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    print('سند فصل چهارم ساخته شد:', out_path)


if __name__ == '__main__':
    build()
